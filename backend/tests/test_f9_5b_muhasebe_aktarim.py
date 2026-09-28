"""F9-5b: muhasebe dışa aktarımı — Luca/Mikro/kanonik serileştiriciler + `GET /export`.

Konu: `app/muhasebe/{serilestirici,hedef_luca,hedef_mikro,hedef_kanonik}.py`,
`app/routers/accounting.py` (`disa_aktar`), `ACTION_TYPES["accounting.exported"]`.
Keşif: `docs/f9-5-muhasebe-disa-aktarim-kesif-2026-09-24.md` §3, §6, §7.2.

--- MUTASYON TABLOSU ------------------------------------------------------

  * `PARCA_FIS_TAVANI` 50 → 51                -> LUCA PARÇA adımı KIRMIZI
  * `parcalar`ı SATIR sınırında bölmek          -> FİŞ PARÇAYA DAĞILMAZ KIRMIZI
  * hedefte `fis_anahtarlari_tekil` çağrısını düşürmek -> LUCA/MİKRO AYNI FİŞ NO KIRMIZI
  * `hucre_metni`nin önekini düşürmek           -> FORMÜL adımı KIRMIZI
  * `tutar_virgullu`da `.`yi korumak            -> MİKRO/KANONİK CSV KIRMIZI
  * uçta `veri.reddedilen` kapısını düşürmek    -> 409 adımı KIRMIZI
  * uçta `fis_anahtarlari_tekil`i düşürmek      -> TEKRARLI FİŞ NO adımı KIRMIZI
  * günlüğü üretecin İÇİNE taşımak              -> KOPAN AKIŞ adımı KIRMIZI
  * `fisler.json`a cari VKN'sini koymak         -> ROLDEN BAĞIMSIZ ÖZET KIRMIZI
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import sys
import tempfile
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal as D
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
DONEM = "2026-07"

_CALISMA = Path(tempfile.mkdtemp(prefix="f95b-"))
os.environ["DATABASE_URL"] = "sqlite:///" + (_CALISMA / "f95b.db").as_posix()
os.environ["SUNGUR_DATA_DIR"] = str(_CALISMA)
os.environ["AUTO_MIGRATE"] = "true"
sys.path.insert(0, str(BACKEND))


# ------------------------------------------------------------------ saf ---

def _fis(no: int, gun: int = 1, satir_sayisi: int = 2, aciklama: str = "Satış", belge_no=None):
    from app.muhasebe.fis import Fis, FisSatiri

    tutar = D("100.00")
    satirlar = [FisSatiri("120", tutar * (satir_sayisi - 1), D("0.00"), aciklama)]
    satirlar += [FisSatiri("600", D("0.00"), tutar, aciklama) for _ in range(satir_sayisi - 1)]
    return Fis(
        fis_no=f"SAT-{no}", fis_tarihi=date(2026, 7, gun), belge_tipi="invoice",
        belge_no=belge_no or f"S-{no}", belge_tarihi=date(2026, 7, gun),
        odeme_yontemi="credit", kaynak="SATIS", satirlar=tuple(satirlar),
        cari_ad="Ahmet", cari_vkn="1234567890",
    )


def _vkn(fis):
    return fis.cari_vkn


def _luca_satirlari(bayt: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    sayfa = load_workbook(io.BytesIO(bayt), read_only=True).active
    return [tuple(c for c in satir) for satir in sayfa.iter_rows(values_only=True)]


@pytest.mark.parametrize("adet, parca", [(50, 1), (51, 2), (100, 2), (101, 3)])
def test_LUCA_PARCA_50_fis(adet, parca) -> None:
    from app.muhasebe.hedef_luca import BASLIKLAR, LucaSerilestirici

    fisler = [_fis(i) for i in range(1, adet + 1)]
    dosyalar = list(LucaSerilestirici().dosyalar(fisler, DONEM, _vkn))
    assert [ad for ad, _ in dosyalar] == [f"luca-{DONEM}-{i:03d}.xlsx" for i in range(1, parca + 1)]
    goruldu = []
    for _, bayt in dosyalar:
        satirlar = _luca_satirlari(bayt)
        assert satirlar[0] == BASLIKLAR
        fis_nolari = list(dict.fromkeys(s[0] for s in satirlar[1:]))
        assert len(fis_nolari) <= 50
        goruldu += fis_nolari
    assert goruldu == [f"SAT-{i}" for i in range(1, adet + 1)]
    if adet >= 51:
        assert "SAT-51" in {s[0] for s in _luca_satirlari(dosyalar[1][1])[1:]}


def test_LUCA_FIS_PARCAYA_DAGILMAZ_ve_hucre_tipleri() -> None:
    from app.muhasebe.hedef_luca import LucaSerilestirici

    # 50 fişin her biri 7 satır: satır sınırında bölen bir mutasyon fişi ikiye ayırırdı.
    fisler = [_fis(i, satir_sayisi=7) for i in range(1, 52)]
    birinci, ikinci = [b for _, b in LucaSerilestirici().dosyalar(fisler, DONEM, _vkn)]
    ilk = _luca_satirlari(birinci)[1:]
    assert len(ilk) == 50 * 7 and {s[0] for s in _luca_satirlari(ikinci)[1:]} == {"SAT-51"}
    satir = ilk[0]
    assert isinstance(satir[1], datetime) and satir[1].date() == date(2026, 7, 1)
    assert satir[4] == 600 and satir[5] == 0  # tutar hücreleri SAYI
    assert satir[6:] == ("S-1", datetime(2026, 7, 1), "invoice", "credit")


def test_LUCA_ayni_fis_no_ValueError() -> None:
    from app.muhasebe.hedef_luca import LucaSerilestirici

    ikiz = [_fis(1, gun=1), _fis(1, gun=2)]
    with pytest.raises(ValueError, match=r"Aynı fiş numarası iki belgede: SAT-1"):
        list(LucaSerilestirici().dosyalar(ikiz, DONEM, _vkn))


def test_MIKRO_ayni_fis_no_ValueError() -> None:
    from app.muhasebe.hedef_mikro import mikro_csv

    ikiz = [_fis(1, gun=1), _fis(1, gun=2)]
    with pytest.raises(ValueError, match=r"Aynı fiş numarası iki belgede: SAT-1"):
        mikro_csv(ikiz)


def test_MIKRO_basligi_BOM_ayri_borc_alacak_virgul() -> None:
    from app.muhasebe.hedef_mikro import MikroSerilestirici

    [(ad, bayt)] = list(MikroSerilestirici().dosyalar([_fis(7, gun=10)], DONEM, _vkn))
    assert ad == f"mikro-{DONEM}.csv"
    assert bayt.startswith(b"\xef\xbb\xbf")
    metin = bayt.decode("utf-8-sig")
    satirlar = metin.split("\r\n")
    assert satirlar[0] == "Fiş No;Fiş Tarihi;Hesap Kodu;Açıklama;Borç;Alacak"
    assert satirlar[1] == "SAT-7;10.07.2026;120;Satış;100,00;0,00"
    assert satirlar[2] == "SAT-7;10.07.2026;600;Satış;0,00;100,00"
    assert satirlar[3] == ""


def test_TUTAR_BICIMI_sabit() -> None:
    from app.muhasebe.serilestirici import tutar_virgullu

    assert tutar_virgullu(D("1234.50")) == "1234,50"
    assert tutar_virgullu(D("0.00")) == "0,00"
    assert tutar_virgullu(D("1000000.05")) == "1000000,05"


def test_KANONIK_ayni_veri_ayni_bayt_ve_ozet() -> None:
    from app.muhasebe.hedef_kanonik import CSV_BASLIKLARI, icerik_sha256, kanonik_csv, kanonik_json

    fisler = [_fis(1), _fis(2, gun=3)]
    a, b = kanonik_json(fisler), kanonik_json(list(fisler))
    assert a == b and icerik_sha256(a) == icerik_sha256(b) == hashlib.sha256(a).hexdigest()
    govde = json.loads(a)
    assert govde[0]["satirlar"][0]["borc"] == "100.00"  # JSON `.` taşır
    assert "cari" not in govde[0] and "1234567890" not in a.decode()
    assert icerik_sha256(kanonik_json([_fis(1)])) != icerik_sha256(a)
    csv_bayt = kanonik_csv(fisler, lambda f: "***")
    assert csv_bayt.startswith(b"\xef\xbb\xbf")
    okunan = list(csv.reader(io.StringIO(csv_bayt.decode("utf-8-sig")), delimiter=";"))
    assert tuple(okunan[0]) == CSV_BASLIKLARI
    assert okunan[1] == ["SAT-1", "2026-07-01", "120", "Satış", "100,00", "0,00", "***", "SATIS"]


def test_FORMUL_ENJEKSIYONU_metin_hucresinde_onek() -> None:
    from openpyxl import load_workbook

    from app.muhasebe.hedef_luca import LucaSerilestirici
    from app.muhasebe.hedef_mikro import mikro_csv
    from app.muhasebe.serilestirici import hucre_metni

    for kotu in ("=1+1", "+90", "-5", "@SUM(A1)"):
        assert hucre_metni(kotu) == "'" + kotu
    assert hucre_metni("Ahmet") == "Ahmet" and hucre_metni(None) == ""
    fis = _fis(1, aciklama='=HYPERLINK("http://x")', belge_no="=cmd")
    [(_, bayt)] = list(LucaSerilestirici().dosyalar([fis], DONEM, _vkn))
    sayfa = load_workbook(io.BytesIO(bayt)).active
    assert sayfa["D2"].data_type == "s" and sayfa["D2"].value == '\'=HYPERLINK("http://x")'
    assert sayfa["G2"].value == "'=cmd"
    okunan = list(csv.reader(io.StringIO(mikro_csv([fis]).decode("utf-8-sig")), delimiter=";"))
    assert okunan[1][3] == '\'=HYPERLINK("http://x")'


# ------------------------------------------------------------ davranış ---

@pytest.fixture(scope="module")
def uygulama():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        yield c


_SAYAC = iter(range(1, 10_000))
PAROLA = "F95bTest!12345"
ROLLER = ("admin", "muhasebe", "satis", "depo", "rapor")
UTC = timezone.utc


def _ekle(db, tablo: str, **alan) -> int:
    from sqlalchemy import text

    sutunlar = ",".join(alan)
    yer = ",".join(":" + k for k in alan)
    return int(db.execute(
        text(f"INSERT INTO {tablo}({sutunlar}) VALUES({yer}) RETURNING id"), alan
    ).scalar_one())


def satis(db, cid, musteri, tarih, kalemler, *, no, durum="approved") -> int:
    """Kalemler `(oran, matrah, kdv)`; kalemsiz onaylı satış DENGESİZ fiştir."""
    kdv = sum((D(k[2]) for k in kalemler), D("0"))
    toplam = sum((D(k[1]) + D(k[2]) for k in kalemler), D("0"))
    oid = _ekle(db, "orders", company_id=cid, customer_id=musteri, order_date=tarih, status=durum,
                document_no=no, payment_method="credit", vat_total=kdv, final_total=toplam,
                subtotal=toplam - kdv, grand_total=toplam)
    for oran, matrah, kv in kalemler:
        tp = D(matrah) + D(kv)
        _ekle(db, "order_items", company_id=cid, order_id=oid, product_name="Ürün",
              quantity=D("1"), unit_price=tp, vat_rate=D(oran), line_subtotal=D(matrah),
              line_vat=D(kv), line_total=tp)
    return oid


def musteri_ekle(db, cid, ad="Ahmet Çiftçi", vkn="1234567890") -> int:
    return _ekle(db, "customers", name=ad, tax_number=vkn, is_active=True, company_id=cid,
                 opening_balance=0, risk_limit=0, payment_term_days=0)


@pytest.fixture()
def dunya(uygulama):
    from app.auth import hash_password
    from app.db import SessionLocal

    n = next(_SAYAC)
    with SessionLocal() as db:
        an = datetime.now(UTC)
        firma = _ekle(db, "companies", name=f"F95b {n}", is_active=True, created_at=an)
        bos_firma = _ekle(db, "companies", name=f"F95b Bos {n}", is_active=True, created_at=an)
        roller = {}
        for rol in ROLLER:
            roller[rol] = _ekle(
                db, "app_users", username=f"f95b-{rol}-{n}", email=f"f95b-{rol}-{n}@f95b.invalid",
                email_verified=True, display_name=rol, password_hash=hash_password(PAROLA), role=rol,
                is_active=True, must_change_password=False, created_at=an,
            )
            for cid in (firma, bos_firma):
                _ekle(db, "user_company_memberships", user_id=roller[rol], company_id=cid,
                      is_default=False, created_at=an)
        musteri = musteri_ekle(db, firma)
        belgeler = [
            satis(db, firma, musteri, "2026-07-10", [("20", "100.00", "20.00")], no="S-1"),
            satis(db, firma, musteri, "2026-07-11", [("10", "50.00", "5.00")], no="=KOTU"),
            satis(db, firma, musteri, "2026-08-01", [("20", "9.00", "1.80")], no="S-AGU"),
        ]
        db.commit()
    yield {"firma": firma, "bos_firma": bos_firma, "n": n, "musteri": musteri, "belgeler": belgeler}


def _baslik(istemci, dunya, rol, cid=None):
    r = istemci.post("/api/auth/login",
                     json={"username": f"f95b-{rol}-{dunya['n']}", "password": PAROLA})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"],
            "X-Company-ID": str(cid or dunya["firma"])}


def _aktar(istemci, dunya, rol="admin", target="luca", period=DONEM, cid=None, ek=""):
    return istemci.get(f"/api/accounting/export?period={period}&target={target}{ek}",
                       headers=_baslik(istemci, dunya, rol, cid))


def _zip(cevap) -> zipfile.ZipFile:
    assert cevap.status_code == 200, cevap.text
    return zipfile.ZipFile(io.BytesIO(cevap.content))


def _gunluk(cid) -> list[dict]:
    from sqlalchemy import text

    from app.db import SessionLocal

    with SessionLocal() as db:
        satirlar = db.execute(text(
            "SELECT details, resource_type, resource_id FROM activity_logs "
            "WHERE company_id=:c AND action_type='accounting.exported' ORDER BY id"),
            {"c": cid}).mappings().all()
    return [
        {**(json.loads(s["details"]) if isinstance(s["details"], str) else s["details"]),
         "_tip": s["resource_type"], "_kimlik": s["resource_id"]}
        for s in satirlar
    ]


def test_LUCA_ZIP_gidis_donus_ve_manifest(uygulama, dunya) -> None:
    r = _aktar(uygulama, dunya)
    assert r.headers["content-type"] == "application/zip"
    assert r.headers["content-disposition"] == f'attachment; filename="muhasebe-{DONEM}-luca.zip"'
    zf = _zip(r)
    assert zf.namelist() == [f"luca-{DONEM}-001.xlsx", "fisler.json", "fisler.csv", "manifest.json"]
    manifest = json.loads(zf.read("manifest.json"))
    fisler = json.loads(zf.read("fisler.json"))
    assert [f["fis_no"] for f in fisler] == [f"SAT-{b}" for b in dunya["belgeler"][:2]]
    assert manifest["icerik_sha256"] == hashlib.sha256(zf.read("fisler.json")).hexdigest()
    assert {k: manifest[k] for k in ("period", "target", "fis_sayisi", "borc_toplami",
                                     "alacak_toplami", "parca_sayisi")} == {
        "period": DONEM, "target": "luca", "fis_sayisi": 2, "borc_toplami": "175.00",
        "alacak_toplami": "175.00", "parca_sayisi": 1,
    }
    assert datetime.fromisoformat(manifest["uretim_zamani"]).utcoffset().total_seconds() == 0
    luca = _luca_satirlari(zf.read(f"luca-{DONEM}-001.xlsx"))
    assert {s[6] for s in luca[1:]} == {"S-1", "'=KOTU"}


@pytest.mark.parametrize("target, beklenen", [
    ("mikro", [f"mikro-{DONEM}.csv", "fisler.json", "fisler.csv", "manifest.json"]),
    ("canonical", ["fisler.json", "fisler.csv", "manifest.json"]),
    ("kanonik", ["fisler.json", "fisler.csv", "manifest.json"]),
])
def test_MIKRO_ve_KANONIK_hedefleri(uygulama, dunya, target, beklenen) -> None:
    zf = _zip(_aktar(uygulama, dunya, target=target))
    assert zf.namelist() == beklenen
    manifest = json.loads(zf.read("manifest.json"))
    assert manifest["target"] == ("canonical" if target == "kanonik" else target)
    assert manifest["parca_sayisi"] == (1 if target == "mikro" else 0)


def test_AYNI_VERI_AYNI_OZET_ve_ROLDEN_BAGIMSIZ(uygulama, dunya) -> None:
    ozetler = {
        json.loads(_zip(_aktar(uygulama, dunya, rol=rol, target=t)).read("manifest.json"))[
            "icerik_sha256"]
        for rol in ("admin", "rapor") for t in ("luca", "mikro", "canonical")
    }
    assert len(ozetler) == 1


def test_VKN_MASKESI_rapor_CSVsinde(uygulama, dunya) -> None:
    def vkn(rol):
        metin = _zip(_aktar(uygulama, dunya, rol=rol, target="canonical")).read("fisler.csv")
        satirlar = list(csv.DictReader(io.StringIO(metin.decode("utf-8-sig")), delimiter=";"))
        return {s["cari_vkn"] for s in satirlar}

    assert vkn("admin") == {"1234567890"}
    [maskeli] = vkn("rapor")
    assert maskeli != "1234567890" and maskeli.endswith("890")


def test_IZIN_rapor_200_satis_depo_403(uygulama, dunya) -> None:
    beklenen = {"admin": 200, "muhasebe": 200, "rapor": 200, "satis": 403, "depo": 403}
    for rol, durum in beklenen.items():
        assert _aktar(uygulama, dunya, rol=rol).status_code == durum, rol


def test_422_logo_bicim_donem(uygulama, dunya) -> None:
    r = _aktar(uygulama, dunya, target="logo")
    assert r.status_code == 422 and r.json()["detail"]["code"] == "HEDEF_DESTEKLENMIYOR"
    assert "Logo" in r.json()["detail"]["message"]
    for target, bicim in (("luca", "csv"), ("mikro", "xlsx"), ("canonical", "xlsx")):
        r = _aktar(uygulama, dunya, target=target, ek=f"&format={bicim}")
        assert r.status_code == 422 and r.json()["detail"]["code"] == "BICIM_DESTEKLENMIYOR"
    for kotu in ("2026-13", "abcd-ef", "2026/07"):
        r = _aktar(uygulama, dunya, period=kotu)
        assert r.status_code == 422 and r.json()["detail"]["code"] == "DONEM_GECERSIZ"
    assert _aktar(uygulama, dunya, period="2026-7").status_code == 422
    r_sap = _aktar(uygulama, dunya, target="sap")
    assert r_sap.status_code == 422 and r_sap.json()["detail"]["code"] == "HEDEF_DESTEKLENMIYOR"
    assert "sap" in r_sap.json()["detail"]["message"]
    assert _gunluk(dunya["firma"]) == []
    assert _aktar(uygulama, dunya, target="luca", ek="&format=xlsx").status_code == 200
    assert _aktar(uygulama, dunya, target="mikro", ek="&format=csv").status_code == 200


def test_BOS_DONEM_200_sifir_fis(uygulama, dunya) -> None:
    zf = _zip(_aktar(uygulama, dunya, cid=dunya["bos_firma"]))
    manifest = json.loads(zf.read("manifest.json"))
    assert manifest["fis_sayisi"] == 0 and manifest["parca_sayisi"] == 0
    assert manifest["borc_toplami"] == "0.00"
    assert json.loads(zf.read("fisler.json")) == []
    assert zf.namelist() == ["fisler.json", "fisler.csv", "manifest.json"]


def test_GUNLUK_aktarim_basina_TEK_satir(uygulama, dunya) -> None:
    zf = _zip(_aktar(uygulama, dunya, target="mikro"))
    [satir] = _gunluk(dunya["firma"])
    manifest = json.loads(zf.read("manifest.json"))
    assert satir == {
        "period": DONEM, "target": "mikro", "fis_sayisi": 2, "borc_toplami": "175.00",
        "icerik_sha256": manifest["icerik_sha256"], "_tip": "account_map", "_kimlik": None,
    }
    _zip(_aktar(uygulama, dunya, target="luca"))
    assert len(_gunluk(dunya["firma"])) == 2


def test_DENGESIZ_DONEM_409_akis_ve_gunluk_YOK(uygulama, dunya) -> None:
    """Kalemsiz onaylı satış DENGESİZ fiştir (veriden, kapı yamanmadan)."""
    from app.db import SessionLocal

    with SessionLocal() as db:
        bos = satis(db, dunya["firma"], dunya["musteri"], "2026-07-20", [], no="S-BOS")
        db.commit()
    r = _aktar(uygulama, dunya)
    assert r.status_code == 409, r.text
    assert r.headers["content-type"].startswith("application/json")
    detay = r.json()["detail"]
    assert detay["code"] == "DONEM_DENGESIZ" and detay["dengesiz_sayisi"] == 1
    assert detay["ilk_hatalar"][0].startswith(f"SAT-{bos}:")
    assert _gunluk(dunya["firma"]) == []


def test_KOPAN_AKIS_gunluk_satirini_KAYBETTIRMEZ(uygulama, dunya, monkeypatch) -> None:
    from app.muhasebe import hedef_luca

    def patla(fisler):
        raise RuntimeError("akış ortasında hata")

    monkeypatch.setattr(hedef_luca, "_xlsx", patla)
    with pytest.raises(RuntimeError, match="akış ortasında"):
        _aktar(uygulama, dunya)
    [satir] = _gunluk(dunya["firma"])
    assert satir["target"] == "luca" and satir["fis_sayisi"] == 2


def test_TEKRARLI_FIS_NO_409_akis_ve_gunluk_YOK(uygulama, dunya, monkeypatch) -> None:
    """Tekrarlı fiş numarası akış ve denetim satırından ÖNCE 409 ile reddedilir."""
    from dataclasses import replace

    from app.routers import accounting

    asli = accounting._donem_kesiti

    def sahte_kesit(cid, donem):
        veri = asli(cid, donem)
        if veri.fisler:
            ikiz = replace(veri.fisler[0], belge_no="S-IKIZ")
            veri.fisler = [veri.fisler[0], ikiz] + list(veri.fisler[1:])
        return veri

    monkeypatch.setattr(accounting, "_donem_kesiti", sahte_kesit)
    r = _aktar(uygulama, dunya)
    assert r.status_code == 409, r.text
    assert r.headers["content-type"].startswith("application/json")
    detay = r.json()["detail"]
    assert detay["code"] == "FIS_NO_TEKRARLI"
    assert "Aynı fiş numarası iki belgede" in detay["message"]
    assert _gunluk(dunya["firma"]) == []


def test_ISLEM_SIRASI_kapi_gunluk_akis() -> None:
    """`disa_aktar`da denge ve tekillik kapıları günlükten, günlük `StreamingResponse`tan ÖNCE."""
    import ast

    kaynak = (BACKEND / "app" / "routers" / "accounting.py").read_text(encoding="utf-8")
    fonk = next(d for d in ast.walk(ast.parse(kaynak))
                if isinstance(d, ast.FunctionDef) and d.name == "disa_aktar")
    metin = ast.unparse(fonk)
    assert metin.index("veri.reddedilen") < metin.index("fis_anahtarlari_tekil") \
        < metin.index("log_request_activity") < metin.index("db.commit()") < metin.index("StreamingResponse(")
